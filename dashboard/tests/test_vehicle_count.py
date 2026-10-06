"""Nombre de véhicules impliqués dans les statistiques, les exports, le tableau de bord et l'admin."""
import csv
import io

from django.urls import reverse
from openpyxl import load_workbook
from pypdf import PdfReader

from dashboard.stats import compute_stats
from reports.models import AccidentReport

from .helpers import PersonasTestCase, make_report


class VehicleCountEverywhereTests(PersonasTestCase):
    def setUp(self):
        super().setUp()
        self.a = make_report(status="VERIFIED", vehicle_count=3)
        self.b = make_report(status="VERIFIED", vehicle_count=2)
        self.none = make_report(status="PENDING", vehicle_count=None)
        self.zero = make_report(status="PENDING", vehicle_count=0)
        self.rejected = make_report(status="REJECTED", vehicle_count=50)
        self.login("staff")

    # ------------------------------------------------------------------ statistiques
    def test_somme_ignorant_les_non_renseignes_et_les_rejetes(self):
        stats = compute_stats(AccidentReport.objects.all())
        self.assertEqual(stats["vehicles"], 5)                   # 3 + 2 + 0 ; None ignoré ; rejeté exclu
        self.assertEqual(stats["vehicle_reports"], 3)             # 3, 2 et 0 sont renseignés ; None et rejeté non comptés
        self.assertEqual(stats["accidents"], 4)

    def test_filtre_rejete_inclut_les_vehicules_des_rejetes(self):
        stats = compute_stats(AccidentReport.objects.filter(status="REJECTED"), include_rejected=True)
        self.assertEqual((stats["vehicles"], stats["vehicle_reports"]), (50, 1))

    def test_base_vide(self):
        AccidentReport.objects.all().delete()
        stats = compute_stats(AccidentReport.objects.all())
        self.assertEqual((stats["vehicles"], stats["vehicle_reports"]), (0, 0))

    def test_carte_du_tableau_de_bord(self):
        page = self.client.get(reverse("dashboard:index"))
        self.assertContains(page, 'id="kpi-vehicles">5<')
        self.assertContains(page, "Véhicules impliqués")
        self.assertContains(page, "3 signalements renseignés")

    def test_statistiques_filtrees(self):
        page = self.client.get(reverse("dashboard:statistics"), {"status": "VERIFIED"})
        self.assertContains(page, 'id="kpi-vehicles">5<')
        page = self.client.get(reverse("dashboard:statistics"), {"status": "PENDING"})
        self.assertContains(page, 'id="kpi-vehicles">0<')
        self.assertContains(page, "1 signalement renseigné")

    # ------------------------------------------------------------------ détail et édition
    def test_detail_non_renseigne_ou_valeur(self):
        page = self.client.get(reverse("dashboard:report_detail", args=[self.none.pk]))
        self.assertContains(page, "Véhicules impliqués")
        self.assertContains(page, "Non renseigné")
        self.assertContains(self.client.get(reverse("dashboard:report_detail", args=[self.zero.pk])), "<dd class=\"col-sm-8\">0</dd>")
        self.assertContains(self.client.get(reverse("dashboard:report_detail", args=[self.a.pk])), "<dd class=\"col-sm-8\">3</dd>")

    def edit_data(self, **overrides):
        data = {"accident_type": "OTHER", "accident_date": "2026-09-01", "accident_time": "07:30", "severity": "LOW",
                "vehicle_count": "2", "injured_count": "1", "death_count": "0", "description": "x"}
        data.update(overrides)
        return data

    def test_edition_peut_vider_ou_renseigner(self):
        url = reverse("dashboard:report_edit", args=[self.a.pk])
        self.client.post(url, self.edit_data(vehicle_count=""))
        self.a.refresh_from_db()
        self.assertIsNone(self.a.vehicle_count)
        self.client.post(url, self.edit_data(vehicle_count="7"))
        self.a.refresh_from_db()
        self.assertEqual(self.a.vehicle_count, 7)
        self.assertEqual(self.client.post(url, self.edit_data(vehicle_count="101")).status_code, 200)

    def test_page_d_edition_libelle_et_prerempli(self):
        page = self.client.get(reverse("dashboard:report_edit", args=[self.a.pk]))
        self.assertContains(page, "Nombre approximatif de véhicules impliqués")
        self.assertRegex(page.content.decode(), r'<input[^>]*name="vehicle_count"[^>]*value="3"')

    # ------------------------------------------------------------------ exports
    def test_csv_valeur_et_vide(self):
        rows = list(csv.reader(io.StringIO(self.client.get(reverse("exports:csv")).content.decode("utf-8-sig")), delimiter=";"))
        header = rows[0]
        self.assertIn("Véhicules impliqués (approx.)", header)
        values = {r[0]: r[header.index("Véhicules impliqués (approx.)")] for r in rows[1:]}
        self.assertEqual(values[self.a.reference], "3")
        self.assertEqual(values[self.zero.reference], "0")
        self.assertEqual(values[self.none.reference], "")                  # non renseigné : cellule vide, jamais « None »
        self.assertNotIn("None", self.client.get(reverse("exports:csv")).content.decode("utf-8-sig"))

    def test_excel_cellule_vide_ou_numerique(self):
        ws = load_workbook(io.BytesIO(self.client.get(reverse("exports:xlsx")).content))["Signalements"]
        col = [c.value for c in ws[1]].index("Véhicules impliqués (approx.)") + 1
        by_ref = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=col).value for r in range(2, ws.max_row + 1)}
        self.assertEqual(by_ref[self.a.reference], 3)
        self.assertEqual(by_ref[self.zero.reference], 0)
        self.assertIsNone(by_ref[self.none.reference])

    def test_excel_resume_et_pdf(self):
        wb = load_workbook(io.BytesIO(self.client.get(reverse("exports:xlsx")).content))
        labels = {row[0].value: row[1].value for row in wb["Résumé"].iter_rows() if row[0].value}
        self.assertEqual(labels["Véhicules impliqués (déclarés)"], 5)
        text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(self.client.get(reverse("exports:pdf")).content)).pages)
        self.assertIn("Véhicules impliqués (déclarés)", text)
        self.assertNotIn("None", text)
        self.assertIn("- / 0 / 0", text)                                       # non renseigné affiché « - » dans le tableau

    def test_exports_filtres_sur_les_vehicules(self):
        rows = list(csv.reader(io.StringIO(self.client.get(reverse("exports:csv"), {"status": "VERIFIED"}).content.decode("utf-8-sig")), delimiter=";"))
        self.assertEqual(sorted(r[rows[0].index("Véhicules impliqués (approx.)")] for r in rows[1:]), ["2", "3"])

    # ------------------------------------------------------------------ admin et confidentialité
    def test_admin_accepte_une_valeur_vide(self):
        url = reverse("admin:reports_accidentreport_change", args=[self.a.pk])
        data = {"status": "VERIFIED", "accident_type": "OTHER", "accident_date": "2026-09-01", "accident_time": "07:30",
                "severity": "LOW", "vehicle_count": "", "injured_count": "1", "death_count": "0", "description": "x",
                "location": "SRID=4326;POINT (1.2231 6.1319)", "admin_note": ""}
        self.assertEqual(self.client.post(url, data).status_code, 302)
        self.a.refresh_from_db()
        self.assertIsNone(self.a.vehicle_count)

    def test_jamais_expose_publiquement(self):
        body = self.client.get(reverse("maps:public_data")).content.decode()
        self.assertNotIn("vehicle", body)
        self.assertNotIn("véhicule", body.lower())
