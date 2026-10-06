"""
Vérification de bout en bout des EXPORTS (CSV, Excel, PDF) et du DJANGO ADMIN dans un vrai navigateur
(Chromium via Playwright), en format smartphone puis ordinateur.

Vérifie : téléchargements réels depuis l'interface, respect des filtres actifs, absence de tout numéro de téléphone,
neutralisation des formules, permissions (utilisateur / staff / superutilisateur), Django Admin (actions, accès).

Prérequis : pip install -r requirements-dev.txt && playwright install chromium
Utilisation : OTP_RESEND_DELAY_SECONDS=0 python manage.py runserver 8000
              E2E_VIEWPORT=mobile|desktop python scripts/e2e_exports_admin_flow.py
(les profils se reconnectent plusieurs fois : le délai anti-renvoi de 60 s doit être désactivé pour CE serveur de test)
Toutes les données de test sont supprimées à la fin.
"""
import csv
import io
import os
import random
import re
import sys
import traceback
from datetime import timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings")
os.environ.setdefault("DJANGO_ALLOW_ASYNC_UNSAFE", "true")
import django  # noqa: E402

django.setup()

from django.contrib.admin.models import LogEntry  # noqa: E402
from django.contrib.gis.geos import Point  # noqa: E402
from django.utils import timezone  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from PIL import Image  # noqa: E402
from playwright.sync_api import sync_playwright  # noqa: E402
from pypdf import PdfReader  # noqa: E402

from accounts.models import OTPCode, User  # noqa: E402
from accounts.roles import ROLE_STAFF, ROLE_SUPERUSER, set_role  # noqa: E402
from reports.choices import AccidentType, Severity  # noqa: E402
from reports.models import AccidentReport  # noqa: E402

BASE = os.environ.get("ACCIMAP_URL", "http://localhost:8000")
SHOTS = os.environ.get("E2E_SHOTS", "e2e_screenshots")
os.makedirs(SHOTS, exist_ok=True)
FORMULA = '=HYPERLINK("http://evil.example","clic")'
results, created_ids, created_users = [], [], []
P = {}


def check(label, cond, detail=""):
    results.append(bool(cond))
    print(("  PASS " if cond else "  FAIL ") + label + (f"  [{detail}]" if detail and not cond else ""))


def phone():
    return "9" + "".join(random.choice("0123456789") for _ in range(7))


def tile_png():
    buffer = io.BytesIO()
    Image.new("RGB", (256, 256), (229, 227, 223)).save(buffer, "PNG")
    return buffer.getvalue()


TILE = tile_png()


def seed():
    assert AccidentReport.objects.count() == 0, "La base de développement doit être vide pour ce test."
    random.seed(5)
    today = timezone.localdate()
    for name, role in [("citizen", None), ("reporter", None), ("staff", ROLE_STAFF), ("super", ROLE_SUPERUSER)]:
        user = User.objects.create_user(phone())
        if role:
            set_role(user, role)
            user.set_password(ADMIN_PASSWORD)
            user.save()
        P[name] = User.objects.get(pk=user.pk)
        created_users.append(P[name])
    types, severities = list(AccidentType.values), list(Severity.values)
    spots = [(1.2228, 6.1319), (1.2050, 6.1650), (1.3000, 6.1250)]

    def make(lon, lat, status, **extra):
        data = dict(accident_type=random.choice(types), severity=random.choice(severities),
                    accident_date=today - timedelta(days=random.randint(1, 60)), accident_time="10:30",
                    injured_count=random.randint(0, 3), death_count=0, location=Point(lon, lat, srid=4326),
                    status=status, is_demo=True, description="Description de test")
        data.update(extra)
        report = AccidentReport.objects.create(**data)
        created_ids.append(report.pk)
        return report

    for status, n in (("PENDING", 6), ("VERIFIED", 10), ("REJECTED", 3)):
        for _ in range(n):
            lon, lat = random.choice(spots)
            anonymous = random.random() < 0.45
            make(random.gauss(lon, 0.003), random.gauss(lat, 0.003), status, is_anonymous=anonymous,
                 user=None if anonymous else P["reporter"])
    for _ in range(2):
        make(1.55, 6.25, "PENDING", user=P["reporter"])           # hors zone indicative
    P["formula"] = make(1.2231, 6.1319, "VERIFIED", is_anonymous=True, description=FORMULA)


def new_page(p, viewport, mobile, accept_downloads=True):
    browser = p.chromium.launch()
    ctx = browser.new_context(viewport=viewport, has_touch=mobile, is_mobile=mobile, locale="fr-FR",
                              device_scale_factor=2 if mobile else 1, accept_downloads=accept_downloads)
    ctx.route("**/tile.openstreetmap.org/**", lambda r: r.fulfill(status=200, content_type="image/png", body=TILE))
    page = ctx.new_page()
    page.set_default_timeout(12000)
    page.errors, page.bad = [], []
    page.on("pageerror", lambda e: page.errors.append(str(e)))
    page.on("console", lambda m: page.errors.append(f"{m.text} @ {m.location.get('url')}")
            if m.type == "error" and "403" not in m.text and "net::ERR" not in m.text else None)
    page.on("response", lambda r: page.bad.append(f"{r.status} {r.url}") if r.status >= 400 and "/static/" in r.url else None)
    return browser, page


def login(page, user, next_path="/dashboard/exports/"):
    page.goto(f"{BASE}/login/?next={next_path}")
    page.fill("#id_phone_number", user.phone_number[4:])
    page.click("button[type=submit]")
    page.wait_for_url(re.compile(r"/verify-otp/"))
    page.fill(".otp-input", page.locator("#dev-code").inner_text().strip())
    page.click("button[type=submit]")
    page.wait_for_load_state("networkidle")

ADMIN_PASSWORD = "E2E-Admin-Phrase-2026"


def admin_login(page, user, next_path="/dashboard/exports/"):
    """Connexion par l'ESPACE PRIVÉ : identifiant administrateur + mot de passe (jamais par l'OTP citoyen)."""
    page.goto(f"{BASE}/dashboard/login/?next={next_path}")
    page.fill("#id_username", user.phone_number[4:])
    page.fill("#id_password", ADMIN_PASSWORD)
    page.click("#admin-login-form button[type=submit]")
    page.wait_for_load_state("networkidle")



def all_phone_forms():
    return [f for u in created_users for f in (u.phone_number, u.phone_number[4:])]


def has_phone(text):
    return any(f in text for f in all_phone_forms()) or re.search(r"\+228\d{8}", text)


def download(page, selector):
    with page.expect_download() as info:
        page.click(selector)
    download = info.value
    path = download.path()
    return download.suggested_filename, Path(path).read_bytes()


def no_overflow(page, label):
    w = page.evaluate("[document.documentElement.scrollWidth, window.innerWidth]")
    check(f"{label} : pas de défilement horizontal", w[0] <= w[1], str(w))


def csv_refs(content):
    rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig")), delimiter=";"))
    return rows[0], rows[1:]


def scenario_exports(p, name, viewport, mobile):
    print(f"\n--- [{name}] EXPORTS depuis l'interface (staff)")
    browser, page = new_page(p, viewport, mobile)
    admin_login(page, P["staff"])
    total = AccidentReport.objects.count()
    check("Page d'export : nombre de signalements = base", page.locator("#export-count strong").inner_text() == str(total))
    check("Page d'export : mention « aucun numéro de téléphone »", "Aucun numéro de téléphone" in page.locator("main").inner_text())
    check("Onglet « Exports » actif", "active" in page.locator(".dash-nav a:has-text('Exports')").get_attribute("class"))
    no_overflow(page, "Page d'export")
    sizes = page.evaluate("[...document.querySelectorAll('main .btn')].filter(e => e.offsetParent).map(e => Math.round(e.getBoundingClientRect().height))")
    check("Page d'export : cibles tactiles >= 44 px", sizes and all(s >= 44 for s in sizes), str(sizes))
    page.screenshot(path=f"{SHOTS}/{name}_export_1_page.png", full_page=True)

    # --- Sans filtre : les trois formats
    filename, data = download(page, "#download-csv")
    headers, rows = csv_refs(data)
    check("CSV : nom de fichier et BOM UTF-8", re.fullmatch(r"accimap_signalements_\d{4}-\d{2}-\d{2}\.csv", filename) and data.startswith(b"\xef\xbb\xbf"))
    check("CSV : toutes les lignes de la base", len(rows) == total and {r[0] for r in rows} == set(AccidentReport.objects.values_list("reference", flat=True)))
    check("CSV : aucun numéro de téléphone", not has_phone(data.decode("utf-8-sig")))
    desc = {r[0]: dict(zip(headers, r))["Description"] for r in rows}
    check("CSV : formule neutralisée (apostrophe)", desc[P["formula"].reference] == "'" + FORMULA)
    filename, data = download(page, "#download-xlsx")
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Signalements"]
    check("Excel : feuilles « Signalements » et « Résumé », bon nombre de lignes", wb.sheetnames == ["Signalements", "Résumé"] and ws.max_row == total + 1)
    check("Excel : en-tête, filtre automatique, volets figés", ws["A1"].value == "Référence" and ws.auto_filter.ref and ws.freeze_panes == "B2")
    check("Excel : la formule n'est jamais une formule", all(c.data_type != "f" for row in ws.iter_rows() for c in row))
    check("Excel : aucun numéro de téléphone", not has_phone("\n".join(str(c.value) for sh in wb for row in sh.iter_rows() for c in row if c.value is not None)))
    filename, data = download(page, "#download-pdf")
    text = "\n".join(pg.extract_text() for pg in PdfReader(io.BytesIO(data)).pages)
    check("PDF : valide, titre ACCIMAP, résumé, aucun filtre", data.startswith(b"%PDF") and "ACCIMAP" in text and "Résumé" in text and "Aucun filtre" in text)
    check("PDF : aucun numéro de téléphone", not has_phone(text))

    # --- Avec filtres : statut Vérifié + mode anonyme (formulaire de la page)
    if mobile and not page.locator("#filters-panel").is_visible():
        page.click("button[data-bs-target='#filters-panel']"); page.wait_for_selector("#filters-panel.show"); page.wait_for_timeout(400)
    page.select_option("#id_status", "VERIFIED"); page.select_option("#id_mode", "anonymous")
    page.click("#export-filters button[type=submit]"); page.wait_for_load_state("networkidle")
    expected = AccidentReport.objects.filter(status="VERIFIED", is_anonymous=True)
    check("Filtres : nombre à exporter = base", page.locator("#export-count strong").inner_text() == str(expected.count()))
    check("Filtres : résumé des filtres affiché", "Statut : Vérifié" in page.locator("#export-filters-summary").inner_text() and "Mode : Anonymes" in page.locator("#export-filters-summary").inner_text())
    _, data = download(page, "#download-csv")
    _, rows = csv_refs(data)
    check("CSV filtré : exactement les signalements vérifiés anonymes", {r[0] for r in rows} == set(expected.values_list("reference", flat=True)))
    _, data = download(page, "#download-xlsx")
    ws = load_workbook(io.BytesIO(data))["Signalements"]
    check("Excel filtré : mêmes lignes", {r[0].value for r in ws.iter_rows(min_row=2)} == set(expected.values_list("reference", flat=True)))
    _, data = download(page, "#download-pdf")
    text = "\n".join(pg.extract_text() for pg in PdfReader(io.BytesIO(data)).pages)
    check("PDF filtré : filtres affichés et références correctes", "Statut : Vérifié" in text and "Mode : Anonymes" in text and all(r in re.sub(r"\s+", "", text) for r in expected.values_list("reference", flat=True)) and f"Signalements ({expected.count()})" in text)
    other = AccidentReport.objects.exclude(pk__in=expected.values_list("pk", flat=True)).first()
    check("PDF filtré : un signalement hors filtre est absent", other.reference not in re.sub(r"\s+", "", text))

    # --- Depuis la liste : les filtres suivent jusqu'à l'export
    page.goto(f"{BASE}/dashboard/reports/?zone=outside"); page.wait_for_load_state("networkidle")
    page.click("#export-link"); page.wait_for_load_state("networkidle")
    check("Liste -> export : le filtre « hors zone » est conservé", page.locator("#export-count strong").inner_text() == "2" and "Zone" in page.locator("#export-filters-summary").inner_text())
    _, data = download(page, "#download-csv")
    _, rows = csv_refs(data)
    check("CSV hors zone : 2 lignes, toutes « Hors zone »", len(rows) == 2 and all(r[9] == "Hors zone" for r in rows))

    # --- Aucun résultat et filtre invalide
    page.goto(f"{BASE}/dashboard/exports/?q=introuvable"); page.wait_for_load_state("networkidle")
    check("Aucun résultat : message clair, aucun bouton de téléchargement", "rien à exporter" in page.locator("main").inner_text() and page.locator("#download-csv").count() == 0)
    page.goto(f"{BASE}/dashboard/exports/?status=XX"); page.wait_for_load_state("networkidle")
    check("Filtre invalide : erreur explicite", "Filtres invalides" in page.locator("main").inner_text())
    entries = LogEntry.objects.filter(change_message__startswith="[EXPORT]", user=P["staff"])
    check("Chaque téléchargement est journalisé (qui, format, filtres) : 7 exports", entries.count() == 7 and not has_phone(" ".join(e.change_message for e in entries)))
    check("Aucune erreur JavaScript", not page.errors, str(page.errors))
    check("Aucun fichier statique manquant", not page.bad, str(page.bad))
    browser.close()


def scenario_export_permissions(p, name, viewport, mobile):
    print(f"\n--- [{name}] EXPORTS : permissions (visiteur, utilisateur)")
    browser, page = new_page(p, viewport, mobile)
    for path in ("/dashboard/exports/", "/dashboard/exports/csv/"):
        r = page.goto(BASE + path)
        check(f"Visiteur -> {path} : redirigé vers la connexion", "/login/" in page.url)
    login(page, P["citizen"], "/")
    for path in ("/dashboard/exports/", "/dashboard/exports/csv/", "/dashboard/exports/xlsx/", "/dashboard/exports/pdf/"):
        status = page.evaluate(f"fetch('{path}').then(r => r.status)")
        check(f"Utilisateur simple -> {path} : 403", status == 403)
    body = page.evaluate("fetch('/dashboard/exports/csv/').then(r => r.text())")
    check("Utilisateur simple : aucune donnée dans la réponse refusée", AccidentReport.objects.first().reference not in body)
    browser.close()


def scenario_django_admin(p, name, viewport, mobile):
    print(f"\n--- [{name}] DJANGO ADMIN (staff puis superutilisateur)")
    browser, page = new_page(p, viewport, mobile)
    admin_login(page, P["staff"], "/admin/")
    check("Staff : accès au Django Admin", "Administration ACCIMAP" in page.content())
    check("Staff : numéro masqué dans l'en-tête", re.search(r"\+228\*{4}\d{2}", page.locator("#user-tools").inner_text()) is not None and not has_phone(page.content()))
    check("Staff : lien de retour vers le tableau de bord", page.locator("a:has-text('Retour au tableau de bord ACCIMAP')").count() == 1)
    check("Staff : pas d'accès aux utilisateurs ni au journal", page.goto(BASE + "/admin/accounts/user/").status == 403 and page.goto(BASE + "/admin/admin/logentry/").status == 403)

    page.goto(BASE + "/admin/reports/accidentreport/"); page.wait_for_load_state("networkidle")
    check("Liste admin : signalements, aucun numéro de téléphone", page.locator("#result_list tbody tr").count() > 0 and not has_phone(page.content()))
    check("Liste admin : filtres (statut, gravité, zone…) présents", all(t in page.locator("#changelist-filter").inner_text() for t in ("Par statut", "Par gravité", "Par zone de couverture")))
    page.screenshot(path=f"{SHOTS}/{name}_admin_1_list.png")

    # Action « Marquer comme vérifié » sur deux signalements en attente
    pending = list(AccidentReport.objects.filter(status="PENDING", is_anonymous=True)[:2])
    page.goto(BASE + "/admin/reports/accidentreport/?status__exact=PENDING"); page.wait_for_load_state("networkidle")
    for report in pending:
        page.check(f"input.action-select[value='{report.pk}']")
    page.select_option("select[name=action]", "mark_verified"); page.click("button[name=index]")
    page.wait_for_selector(".messagelist .success")
    for report in pending:
        report.refresh_from_db()
    check("Action admin « Marquer comme vérifié » : statut et traitement enregistrés", all(r.status == "VERIFIED" and r.verified_by_id == P["staff"].pk for r in pending))

    # Page de modification : relations utilisateur absentes, photo jamais via /media/
    identified = AccidentReport.objects.filter(is_anonymous=False).first()
    page.goto(f"{BASE}/admin/reports/accidentreport/{identified.pk}/change/"); page.wait_for_load_state("networkidle")
    html = page.content()
    check("Fiche admin : déclarant masqué, aucun champ « utilisateur »", "Identifié (+228****" in html and 'name="user"' not in html and 'name="verified_by"' not in html and not has_phone(html))
    page.screenshot(path=f"{SHOTS}/{name}_admin_2_change.png", full_page=True)
    r = page.goto(f"{BASE}/admin/reports/accidentreport/{identified.pk}/history/")
    check("Historique admin : utilisateurs masqués", r.status == 200 and not has_phone(page.content()))
    # Le widget de carte du Django Admin charge OpenLayers depuis un CDN : sans accès internet (sandbox), seule
    # l'erreur « ol is not defined » est attendue. Toute autre erreur reste un échec.
    others = [e for e in page.errors if "ol is not defined" not in e]
    check("Aucune erreur JavaScript côté staff (hors CDN OpenLayers de la carte d'admin)", not others, str(others))
    browser.close()

    browser, page = new_page(p, viewport, mobile)
    admin_login(page, P["super"], "/admin/")
    page.goto(BASE + "/admin/accounts/user/"); page.wait_for_load_state("networkidle")
    check("Superutilisateur : liste des utilisateurs avec numéros complets", P["staff"].phone_number in page.content())
    page.screenshot(path=f"{SHOTS}/{name}_admin_3_users.png")
    page.goto(BASE + "/admin/admin/logentry/"); page.wait_for_load_state("networkidle")
    check("Superutilisateur : journal des actions, exports tracés, numéros masqués", "Export" in page.content() and "+228****" in page.content() and not any(f in page.content() for f in all_phone_forms()))
    check("Journal en lecture seule (ajout interdit)", page.goto(BASE + "/admin/admin/logentry/add/").status == 403)
    browser.close()


def cleanup():
    reports = AccidentReport.objects.filter(pk__in=created_ids)
    n = reports.delete()[0]
    LogEntry.objects.filter(user__in=[u.pk for u in created_users]).delete()
    OTPCode.objects.filter(phone_number__in=[u.phone_number for u in created_users]).delete()
    User.objects.filter(pk__in=[u.pk for u in created_users]).delete()
    print(f"\nNettoyage : {n} objet(s) et {len(created_users)} utilisateur(s) d'essai supprimés.")


if __name__ == "__main__":
    which = os.environ.get("E2E_VIEWPORT", "both")
    suites = [("mobile", {"width": 390, "height": 844}, True), ("desktop", {"width": 1366, "height": 800}, False)]
    suites = [s for s in suites if which in ("both", s[0])]
    seed()
    print(f"Jeu de test : {AccidentReport.objects.count()} signalements, {len(created_users)} comptes.")
    try:
        with sync_playwright() as p:
            for name, viewport, mobile in suites:
                print(f"\n================ {name.upper()} {viewport['width']}x{viewport['height']} ================")
                for scenario in (scenario_exports, scenario_export_permissions, scenario_django_admin):
                    try:
                        scenario(p, name, viewport, mobile)
                    except Exception:
                        check(f"{scenario.__name__} s'est exécuté sans exception", False, traceback.format_exc(limit=3))
    finally:
        cleanup()
    ok = sum(results)
    print(f"\nRÉSULTAT : {ok}/{len(results)} contrôles réussis")
    sys.exit(0 if ok == len(results) else 1)
