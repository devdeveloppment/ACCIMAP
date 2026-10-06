"""
Vérification de bout en bout du PARCOURS COMPLET ACCIMAP (phase 8) dans un vrai navigateur (Chromium via Playwright),
enchaîné sans raccourci, en format smartphone puis ordinateur :

    1. Citoyen : accueil -> « Signaler » -> connexion OTP (numéro -> code -> vérification)
    2. Signalement en 4 étapes (type, lieu GPS, détails, récapitulatif) -> enregistré « En attente »
    3. Absent de la carte publique tant qu'il n'est pas validé
    4. Administrateur : connexion PRIVÉE (identifiant + mot de passe) -> fiche -> « Vérifié »
    5. Apparition sur la carte publique (compteur, repère, API sans donnée personnelle)
    6. Statistiques mises à jour (indicateurs et graphiques)
    7. Exports Excel et PDF téléchargés depuis l'interface, contenu contrôlé
    8. Permissions : le citoyen n'accède pas à l'espace privé

Prérequis : pip install -r requirements-dev.txt && playwright install chromium
Utilisation : python manage.py runserver 8000   (OTP_DEV_MODE=True)
              E2E_VIEWPORT=mobile|desktop python scripts/e2e_full_journey.py     (défaut : les deux)
Les tuiles OpenStreetMap et le service de quartier sont simulés. Toutes les données créées sont supprimées à la fin.
"""
import io
import json
import os
import random
import re
import sys
import traceback
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings")
os.environ.setdefault("DJANGO_ALLOW_ASYNC_UNSAFE", "true")
import django  # noqa: E402

django.setup()

from openpyxl import load_workbook  # noqa: E402
from PIL import Image  # noqa: E402
from playwright.sync_api import sync_playwright  # noqa: E402
from pypdf import PdfReader  # noqa: E402

from accounts.models import OTPCode, User  # noqa: E402
from accounts.roles import ROLE_SUPERUSER, set_role  # noqa: E402
from reports.models import AccidentReport  # noqa: E402

BASE = os.environ.get("ACCIMAP_URL", "http://localhost:8000")
SHOTS = os.environ.get("E2E_SHOTS", "e2e_screenshots")
os.makedirs(SHOTS, exist_ok=True)

GPS = {"latitude": 6.1402, "longitude": 1.2131}   # Lomé, dans la zone indicative
PLACE = "Bè, Lomé"
ADMIN_PASSWORD = "E2E-Parcours-Complet-2026"
PRIVATE = "Témoin : Afi, 90 00 11 22, plaque TG-7788-AC"
VIEWPORTS = {
    "mobile": ({"width": 390, "height": 844}, True),
    "desktop": ({"width": 1366, "height": 900}, False),
}
results, created_refs, created_phones = [], [], []


def check(label, cond, detail=""):
    results.append(bool(cond))
    print(("  PASS " if cond else "  FAIL ") + label + (f"  [{detail}]" if detail and not cond else ""))
    return bool(cond)


def fresh_phone():
    n = "9" + "".join(random.choice("0123456789") for _ in range(7))
    created_phones.append("+228" + n)
    return n


def tile_png():
    buffer = io.BytesIO()
    Image.new("RGB", (256, 256), (229, 227, 223)).save(buffer, "PNG")
    return buffer.getvalue()


TILE = tile_png()


def new_context(browser, viewport, mobile):
    ctx = browser.new_context(viewport=viewport, has_touch=mobile, is_mobile=mobile, locale="fr-FR",
                              geolocation={**GPS, "accuracy": 12}, permissions=["geolocation"],
                              accept_downloads=True)
    ctx.route("**/tile.openstreetmap.org/**", lambda r: r.fulfill(status=200, content_type="image/png", body=TILE))
    ctx.route("**/report/place/**", lambda r: r.fulfill(status=200, content_type="application/json",
                                                         body=json.dumps({"place": PLACE})))
    page = ctx.new_page()
    page.set_default_timeout(15000)
    page.errors = []
    page.on("pageerror", lambda e: page.errors.append(str(e)))
    page.on("console", lambda m: page.errors.append(f"{m.text} @ {m.location.get('url')}") if m.type == "error" else None)
    page.on("response", lambda r: page.errors.append(f"{r.status} {r.url}")
            if r.status >= 400 and "/static/" in r.url else None)
    return ctx, page


def wait_step(page, n):
    page.wait_for_selector(f".report-step.is-active[data-step='{n}']")
    page.wait_for_timeout(350)


def bump(page, field_id, times=1):
    for _ in range(times):
        page.locator(f".counter:has(#{field_id}) [data-delta='1']").click()


def kpi(page, name):
    return int(page.locator(f"#kpi-{name}").inner_text().strip())


def map_ready(page):
    page.wait_for_function("document.getElementById('map-loading').hidden"
                           " && !/Chargement/.test(document.getElementById('map-count').textContent)")
    page.wait_for_timeout(300)
    m = re.match(r"(\d+)", page.locator("#map-count").inner_text())
    return int(m.group(1)) if m else 0


def download(page, selector):
    with page.expect_download() as info:
        page.click(selector)
    dl = info.value
    return dl.suggested_filename, Path(dl.path()).read_bytes()


def make_admin():
    user = User.objects.create_user(fresh_phone())
    set_role(user, ROLE_SUPERUSER)
    user = User.objects.get(pk=user.pk)
    user.set_password(ADMIN_PASSWORD)
    user.save()
    return user


# --------------------------------------------------------------------------------------------------
def journey(p, name, viewport, mobile):
    print(f"\n=== [{name}] Parcours complet citoyen -> administration -> carte -> statistiques -> exports")
    admin = make_admin()
    browser = p.chromium.launch()
    citizen_ctx, page = new_context(browser, viewport, mobile)

    # ---- Carte publique : compteur de départ
    page.goto(BASE + "/map/")
    public_before = map_ready(page)

    # ---- 1. Citoyen : connexion OTP
    print("-- 1. Authentification OTP")
    page.goto(BASE + "/")
    page.click(".hero a.btn-accent")
    page.wait_for_url(re.compile(r"/login/\?next="))
    check("Signaler sans être connecté -> page de connexion", "/login/" in page.url)
    phone = fresh_phone()
    page.fill("#id_phone_number", phone)
    page.click("button[type=submit]")
    page.wait_for_url(re.compile(r"/verify-otp/"))
    check("Numéro accepté -> saisie du code", "/verify-otp/" in page.url)
    otp = OTPCode.objects.filter(phone_number="+228" + phone).first()
    check("BASE : code OTP créé, stocké sous forme d'empreinte uniquement",
          otp and not otp.is_used and len(otp.code_hash) == 64 and page.locator("#dev-code").inner_text().strip() not in otp.code_hash)
    page.fill(".otp-input", "000000" if page.locator("#dev-code").inner_text().strip() != "000000" else "111111")
    page.click("button[type=submit]")
    page.wait_for_load_state("networkidle")
    check("Code erroné -> refusé, reste sur la vérification", "/verify-otp/" in page.url and "incorrect" in page.inner_text("body"))
    page.fill(".otp-input", page.locator("#dev-code").inner_text().strip())
    page.click("button[type=submit]")
    page.wait_for_url(re.compile(r"/report/"))
    citizen = User.objects.get(phone_number="+228" + phone)
    check("Code correct -> connecté et renvoyé vers le signalement", page.url.split("#")[0].endswith("/report/"))
    check("BASE : compte citoyen créé, sans mot de passe ni accès staff",
          not citizen.is_staff and not citizen.has_usable_password())

    # ---- 2. Signalement en 4 étapes
    print("-- 2. Signalement")
    page.wait_for_load_state("networkidle")
    page.locator("label[for='type-PEDESTRIAN']").click()
    wait_step(page, 2)
    page.wait_for_function("document.getElementById('id_latitude').value !== ''")
    lat, lon = float(page.input_value("#id_latitude")), float(page.input_value("#id_longitude"))
    check("GPS accepté : coordonnées reçues et affichées", abs(lat - GPS["latitude"]) < 1e-4 and abs(lon - GPS["longitude"]) < 1e-4, f"{lat},{lon}")
    check("Repère positionné sur la carte", page.locator("#location-map .leaflet-marker-icon").count() == 1)
    page.wait_for_function(f"document.getElementById('place-label').textContent.includes('{PLACE}')")
    check("Lieu lisible affiché (géocodage inverse)", PLACE in page.locator("#place-label").inner_text())
    page.click("#confirm-position-btn")
    wait_step(page, 3)
    page.fill("#id_description", PRIVATE)
    page.click("label.btn:has-text('Grave') >> nth=0")
    bump(page, "id_vehicle_count", 2)
    bump(page, "id_injured_count", 2)
    page.click("#step3-next")
    wait_step(page, 4)
    page.screenshot(path=f"{SHOTS}/{name}_full_1_recap.png", full_page=True)
    page.click("#submit-report")
    page.wait_for_selector("#report-reference")
    reference = page.locator("#report-reference").inner_text().strip()
    created_refs.append(reference)
    check("Confirmation affichée avec la référence", re.fullmatch(r"ACC-\d{4}-\d{6}", reference), reference)
    report = AccidentReport.objects.get(reference=reference)
    check("BASE : rattaché au citoyen, statut « En attente »", report.user_id == citizen.pk and report.status == "PENDING")
    check("BASE : PointField SRID 4326 = coordonnées GPS reçues (non modifiées)",
          report.location.srid == 4326 and abs(report.location.y - GPS["latitude"]) < 1e-6 and abs(report.location.x - GPS["longitude"]) < 1e-6,
          f"{report.location.y},{report.location.x}")
    check("BASE : détails enregistrés (piéton, grave, 2 véhicules, 2 blessés)",
          (report.accident_type, report.severity, report.vehicle_count, report.injured_count) == ("PEDESTRIAN", "SEVERE", 2, 2))

    # ---- 3. Pas encore public ; le citoyen n'accède pas à l'espace privé
    print("-- 3. Avant validation")
    page.goto(BASE + "/map/")
    check("Carte publique : signalement EN ATTENTE non affiché", map_ready(page) == public_before)
    for path in ("/dashboard/", "/dashboard/statistics/", "/dashboard/exports/"):
        response = page.goto(BASE + path)
        check(f"Citoyen connecté par OTP : {path} refusé (403)", response.status == 403, str(response.status))

    # ---- 4. Administrateur : connexion privée et validation
    print("-- 4. Traitement administratif")
    admin_ctx, admin_page = new_context(browser, viewport, mobile)
    admin_page.goto(BASE + "/dashboard/")
    check("Espace privé sans session -> connexion privée", "/dashboard/login/" in admin_page.url)
    admin_user = admin
    admin_page.fill("#id_username", admin_user.phone_number[4:])
    admin_page.fill("#id_password", "mauvais-mot-de-passe")
    admin_page.click("#admin-login-form button[type=submit]")
    admin_page.wait_for_load_state("networkidle")
    check("Mot de passe erroné -> refusé", "/dashboard/login/" in admin_page.url and "incorrect" in admin_page.inner_text("body"))
    admin_page.fill("#id_username", admin_user.phone_number[4:])
    admin_page.fill("#id_password", ADMIN_PASSWORD)
    admin_page.click("#admin-login-form button[type=submit]")
    admin_page.wait_for_load_state("networkidle")
    check("Identifiant + mot de passe -> tableau de bord", admin_page.url.rstrip("/").endswith("/dashboard"))
    admin_page.goto(BASE + "/dashboard/statistics/")
    admin_page.wait_for_load_state("networkidle")
    before = {k: kpi(admin_page, k) for k in ("total", "pending", "verified", "injured", "identified")}

    admin_page.goto(f"{BASE}/dashboard/reports/?q={reference}")
    admin_page.wait_for_load_state("networkidle")
    check("Liste admin : le signalement est retrouvé par sa référence", reference in admin_page.inner_text("main"))
    admin_page.goto(f"{BASE}/dashboard/reports/{report.pk}/")
    admin_page.wait_for_load_state("networkidle")
    check("Fiche : référence, statut « En attente », mini-carte",
          reference in admin_page.inner_text("main") and "En attente" in admin_page.locator("#status-badge").inner_text()
          and admin_page.locator("#detail-map .leaflet-marker-icon").count() == 1)
    check("Fiche : déclarant masqué (jamais le numéro complet)", phone not in admin_page.content())
    admin_page.fill("#id_admin_note", "Confirmé par la police municipale")
    admin_page.click("button[name=status][value=VERIFIED]")
    admin_page.wait_for_selector(".alert-success")
    check("Validation : message « Statut mis à jour : Vérifié. »", "Statut mis à jour : Vérifié." in admin_page.locator(".alert-success").inner_text())
    report.refresh_from_db()
    check("BASE : Vérifié, par cet administrateur, avec date et note",
          report.status == "VERIFIED" and report.verified_by_id == admin_user.pk and report.verified_at
          and report.admin_note == "Confirmé par la police municipale")
    admin_page.screenshot(path=f"{SHOTS}/{name}_full_2_verified.png", full_page=True)

    # ---- 5. Carte publique
    print("-- 5. Carte publique")
    page.errors.clear()   # les 403 attendus de l'étape 3 ne concernent pas la carte
    page.goto(BASE + "/map/")
    check("Carte publique : compteur +1 après validation", map_ready(page) == public_before + 1)
    check("Carte publique : repère ou groupe affiché",
          page.locator(".leaflet-marker-icon, .marker-cluster").count() > 0)
    api = page.request.get(BASE + "/map/data/")
    body = api.text()
    coords = [f["geometry"]["coordinates"] for f in api.json()["features"]]
    check("API publique : point aux coordonnées enregistrées", [round(GPS["longitude"], 6), round(GPS["latitude"], 6)] in coords)
    check("API publique : aucune donnée personnelle ni interne",
          not any(s in body for s in (phone, "Afi", "TG-7788-AC", "police municipale", reference, str(report.pk))))
    check("Carte publique : aucune erreur JavaScript ni ressource manquante", not page.errors, "; ".join(page.errors[:3]))
    page.screenshot(path=f"{SHOTS}/{name}_full_3_public_map.png")

    # ---- 6. Statistiques
    print("-- 6. Statistiques")
    admin_page.goto(BASE + "/dashboard/statistics/")
    admin_page.wait_for_load_state("networkidle")
    after = {k: kpi(admin_page, k) for k in ("total", "pending", "verified", "injured", "identified")}
    check("Statistiques : Vérifiés +1", after["verified"] == before["verified"] + 1, f"{before} -> {after}")
    check("Statistiques : En attente -1", after["pending"] == before["pending"] - 1, f"{before} -> {after}")
    check("Statistiques : total et blessés inchangés par la validation",
          (after["total"], after["injured"]) == (before["total"], before["injured"]))
    charts = admin_page.evaluate("[...document.querySelectorAll('canvas')].filter(c => c.width > 0 && c.height > 0).length")
    check("Statistiques : graphiques Chart.js dessinés", charts >= 3, str(charts))
    check("Tableau de bord : aucune erreur JavaScript", not admin_page.errors, "; ".join(admin_page.errors[:3]))

    # ---- 7. Exports Excel et PDF depuis l'interface (filtre : Vérifié)
    print("-- 7. Exports")
    admin_page.goto(BASE + "/dashboard/exports/?status=VERIFIED")
    admin_page.wait_for_load_state("networkidle")
    filename, data = download(admin_page, "#download-xlsx")
    sheet = load_workbook(io.BytesIO(data)).active
    rows = list(sheet.iter_rows(values_only=True))
    line = next((dict(zip(rows[0], r)) for r in rows[1:] if r[0] == reference), None)
    check("Excel : fichier .xlsx téléchargé", filename.endswith(".xlsx"), filename)
    check("Excel : ligne du signalement validé, statut, coordonnées et zone exacts",
          line and line["Statut"] == "Vérifié" and abs(line["Latitude"] - GPS["latitude"]) < 1e-6
          and abs(line["Longitude"] - GPS["longitude"]) < 1e-6 and line["Zone de couverture"] == "Dans la zone",
          str(line))
    check("Excel : seulement des signalements vérifiés (filtre respecté)", all(r[5] == "Vérifié" for r in rows[1:] if r[0]))
    flat = json.dumps([[str(c) for c in r] for r in rows])
    check("Excel : aucun numéro de téléphone ni note interne", phone not in flat and "+228" not in flat and "police municipale" not in flat)
    filename, data = download(admin_page, "#download-pdf")
    text = "".join(pg.extract_text() for pg in PdfReader(io.BytesIO(data)).pages)
    check("PDF : fichier .pdf téléchargé et lisible", filename.endswith(".pdf") and data[:4] == b"%PDF", filename)
    check("PDF : contient la référence du signalement", reference in re.sub(r"\s+", "", text))
    check("PDF : aucun numéro de téléphone", phone not in re.sub(r"\s+", "", text) and "+228" not in text)

    # ---- 8. Déconnexions
    citizen_ctx.close(); admin_ctx.close(); browser.close()


def cleanup():
    AccidentReport.objects.filter(reference__in=created_refs).delete()
    User.objects.filter(phone_number__in=created_phones).delete()
    OTPCode.objects.filter(phone_number__in=created_phones).delete()


def main():
    wanted = os.environ.get("E2E_VIEWPORT")
    names = [wanted] if wanted else list(VIEWPORTS)
    try:
        with sync_playwright() as p:
            for name in names:
                viewport, mobile = VIEWPORTS[name]
                try:
                    journey(p, name, viewport, mobile)
                except Exception:
                    traceback.print_exc()
                    check(f"[{name}] parcours exécuté jusqu'au bout", False)
    finally:
        cleanup()
    passed = sum(results)
    print(f"\nRésultat : {passed}/{len(results)} vérifications réussies")
    sys.exit(0 if passed == len(results) else 1)


if __name__ == "__main__":
    main()
