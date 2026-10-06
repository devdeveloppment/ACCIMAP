import csv
import io
import re
from datetime import timedelta

from django.contrib.admin.models import LogEntry
from django.contrib.gis.geos import Point
from django.test import override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from pypdf import PdfReader

from accounts.models import User
from dashboard.tests.helpers import PersonasTestCase, make_report
from exports import services
from reports.models import AccidentReport

PARIS = Point(2.3522, 48.8566, srid=4326)
PHONE_PATTERN = re.compile(r"\+228|(?<!\d)9\d{7}(?!\d)")
FORMULA = '=HYPERLINK("http://evil.example","clic")'


def csv_rows(response):
    text = response.content.decode("utf-8-sig")
    return list(csv.reader(io.StringIO(text), delimiter=";"))


def pdf_text(response):
    return "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(response.content)).pages)


def squash(text):
    """Texte sans espaces ni retours à la ligne (une cellule de tableau peut passer à la ligne)."""
    return re.sub(r"\s+", "", text)


class ExportDataMixin:
    """Jeu de données : 3 statuts, identifiés/anonymes, dans/hors zone, texte piégé."""

    def make_data(self):
        today = timezone.localdate()
        self.owner = User.objects.create_user("90123456")
        self.v = make_report(status="VERIFIED", user=self.owner, accident_type="ROLLOVER", severity="SEVERE",
                             accident_date=today - timedelta(days=10), description="Carrefour du marché",
                             injured_count=2, death_count=1, vehicle_count=3, admin_note="NOTE-INTERNE-CONFIDENTIELLE")
        self.p = make_report(status="PENDING", is_anonymous=True, accident_type="OTHER", severity="LOW",
                             accident_date=today - timedelta(days=3), description=FORMULA)
        self.r = make_report(status="REJECTED", user=self.owner, accident_type="OTHER", severity="CRITICAL",
                             accident_date=today - timedelta(days=1), location=PARIS, description="Hors zone")
        self.refs = {self.v.reference, self.p.reference, self.r.reference}

    def refs_in(self, rows):
        return {row[0] for row in rows[1:]}


class CsvExportTests(ExportDataMixin, PersonasTestCase):
    def setUp(self):
        super().setUp()
        self.make_data()
        self.login("staff")
        self.url = reverse("exports:csv")

    def test_en_tetes_et_format(self):
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "text/csv; charset=utf-8")
        self.assertRegex(response["Content-Disposition"], r'attachment; filename="accimap_signalements_\d{4}-\d{2}-\d{2}\.csv"')
        self.assertTrue(response.content.startswith(b"\xef\xbb\xbf"))     # BOM : accents corrects dans Excel
        rows = csv_rows(response)
        self.assertEqual(rows[0], [c.header for c in services.COLUMNS])
        self.assertEqual(len(rows), 4)
        self.assertEqual(self.refs_in(rows), self.refs)

    def test_contenu_d_une_ligne(self):
        rows = csv_rows(self.client.get(self.url, {"status": "VERIFIED"}))
        row = dict(zip(rows[0], rows[1]))
        self.assertEqual(row["Référence"], self.v.reference)
        self.assertEqual(row["Type d'accident"], "Renversement")
        self.assertEqual((row["Gravité"], row["Statut"], row["Mode"]), ("Grave", "Vérifié", "Identifié"))
        self.assertEqual((row["Latitude"], row["Longitude"]), ("6.172500", "1.231400"))   # point décimal, 6 décimales
        self.assertEqual(row["Zone de couverture"], "Dans la zone")
        self.assertEqual((row["Véhicules impliqués (approx.)"], row["Blessés"], row["Décès"]), ("3", "2", "1"))
        self.assertEqual(row["Description"], "Carrefour du marché")
        self.assertEqual(row["Photo jointe"], "Non")
        self.assertRegex(row["Heure"], r"^\d{2}:\d{2}$")
        self.assertRegex(row["Date de l'accident"], r"^\d{4}-\d{2}-\d{2}$")

    def test_mode_et_zone(self):
        rows = {r[0]: dict(zip(csv_rows(self.client.get(self.url))[0], r)) for r in csv_rows(self.client.get(self.url))[1:]}
        self.assertEqual(rows[self.p.reference]["Mode"], "Anonyme")
        self.assertEqual(rows[self.r.reference]["Zone de couverture"], "Hors zone")

    def test_injection_de_formule_neutralisee(self):
        rows = csv_rows(self.client.get(self.url, {"status": "PENDING"}))
        description = dict(zip(rows[0], rows[1]))["Description"]
        self.assertEqual(description, "'" + FORMULA)

    def test_aucun_numero_de_telephone(self):
        body = self.client.get(self.url).content.decode("utf-8-sig")
        self.assertIsNone(PHONE_PATTERN.search(body))
        for form in self.phone_forms(self.owner):
            self.assertNotIn(form, body)

    def test_note_interne_et_traitement_exclus(self):
        self.v.set_status("VERIFIED", self.personas["staff"], note="NOTE-INTERNE-CONFIDENTIELLE")
        body = self.client.get(self.url).content.decode("utf-8-sig")
        self.assertNotIn("NOTE-INTERNE-CONFIDENTIELLE", body)

    def test_filtres(self):
        cases = [
            ({"status": "VERIFIED"}, {self.v}), ({"status": "REJECTED"}, {self.r}),
            ({"mode": "anonymous"}, {self.p}), ({"mode": "identified"}, {self.v, self.r}),
            ({"accident_type": "OTHER"}, {self.p, self.r}), ({"severity": "SEVERE"}, {self.v}),
            ({"zone": "outside"}, {self.r}), ({"zone": "inside"}, {self.v, self.p}),
            ({"q": "marché"}, {self.v}), ({"q": self.p.reference}, {self.p}),
            ({"status": "OTHER_IGNORED_BY_NOTHING"}, None),
        ]
        for params, expected in cases:
            with self.subTest(params=params):
                response = self.client.get(self.url, params)
                if expected is None:
                    self.assertEqual(response.status_code, 302)   # filtre invalide : pas de fichier
                    continue
                self.assertEqual(self.refs_in(csv_rows(response)), {r.reference for r in expected})

    def test_filtre_periode_et_combinaison(self):
        today = timezone.localdate()
        rows = csv_rows(self.client.get(self.url, {"date_from": (today - timedelta(days=4)).isoformat()}))
        self.assertEqual(self.refs_in(rows), {self.p.reference, self.r.reference})
        rows = csv_rows(self.client.get(self.url, {"date_to": (today - timedelta(days=5)).isoformat()}))
        self.assertEqual(self.refs_in(rows), {self.v.reference})
        rows = csv_rows(self.client.get(self.url, {"accident_type": "OTHER", "mode": "identified", "zone": "outside"}))
        self.assertEqual(self.refs_in(rows), {self.r.reference})

    def test_aucun_resultat_ou_filtre_invalide_redirige_avec_message(self):
        response = self.client.get(self.url, {"q": "introuvable"}, follow=True)
        self.assertRedirects(response, reverse("exports:index") + "?q=introuvable")
        self.assertContains(response, "rien à exporter")
        response = self.client.get(self.url, {"date_from": "2026-05-10", "date_to": "2026-05-01"}, follow=True)
        self.assertContains(response, "Filtres invalides")

    @override_settings(EXPORT_MAX_ROWS=2)
    def test_trop_de_lignes_refuse_au_lieu_de_tronquer_en_silence(self):
        response = self.client.get(self.url, follow=True)
        self.assertContains(response, "trop pour ce format")
        self.assertNotEqual(response["Content-Type"], "text/csv; charset=utf-8")

    def test_la_requete_ne_charge_jamais_l_utilisateur(self):
        sql = str(services.export_queryset(AccidentReport.objects.all()).query)
        self.assertNotIn("user_id", sql)
        self.assertNotIn("verified_by_id", sql)
        self.assertNotIn("admin_note", sql)

    def test_aucune_colonne_d_identite(self):
        for column in services.COLUMNS:
            self.assertNotRegex(column.key + column.header.lower(), r"phone|téléphone|user|utilisateur|déclarant|compte")


class XlsxExportTests(ExportDataMixin, PersonasTestCase):
    def setUp(self):
        super().setUp()
        self.make_data()
        self.login("staff")
        self.url = reverse("exports:xlsx")

    def workbook(self, **params):
        response = self.client.get(self.url, params)
        self.assertEqual(response.status_code, 200)
        return response, load_workbook(io.BytesIO(response.content))

    def test_type_nom_et_feuilles(self):
        response, wb = self.workbook()
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertRegex(response["Content-Disposition"], r'filename="accimap_signalements_[\d-]+\.xlsx"')
        self.assertEqual(wb.sheetnames, ["Signalements", "Résumé"])

    def test_structure_de_la_feuille(self):
        _, wb = self.workbook()
        ws = wb["Signalements"]
        self.assertEqual([c.value for c in ws[1]], [c.header for c in services.COLUMNS])
        self.assertTrue(all(c.font.bold for c in ws[1]))
        self.assertEqual(ws.freeze_panes, "B2")
        self.assertEqual(ws.auto_filter.ref, f"A1:R{ws.max_row}")
        self.assertEqual(ws.max_row, 4)
        self.assertEqual(ws.column_dimensions["N"].width, 50)     # description lisible

    def test_valeurs_typees(self):
        _, wb = self.workbook(status="VERIFIED")
        ws = wb["Signalements"]
        values = {ws.cell(row=1, column=c).value: ws.cell(row=2, column=c) for c in range(1, 19)}
        self.assertEqual(values["Référence"].value, self.v.reference)
        self.assertIsInstance(values["Latitude"].value, float)
        self.assertAlmostEqual(values["Latitude"].value, 6.1725)
        self.assertAlmostEqual(values["Longitude"].value, 1.2314)
        self.assertEqual(values["Latitude"].number_format, "0.000000")
        self.assertEqual(values["Blessés"].value, 2)
        self.assertEqual(values["Date de l'accident"].number_format, "DD/MM/YYYY")
        self.assertEqual(values["Date de l'accident"].value.date(), self.v.accident_date)

    def test_injection_de_formule_neutralisee(self):
        _, wb = self.workbook(status="PENDING")
        cell = wb["Signalements"]["N2"]
        self.assertEqual(cell.value, "'" + FORMULA)
        self.assertEqual(cell.data_type, "s")          # texte, jamais une formule

    def test_caracteres_de_controle_supprimes(self):
        make_report(status="VERIFIED", description="avant\x07\x0bapres")
        _, wb = self.workbook(q="avant")
        self.assertEqual(wb["Signalements"]["N2"].value, "avantapres")

    def test_filtres_respectes(self):
        _, wb = self.workbook(mode="identified", zone="outside")
        ws = wb["Signalements"]
        self.assertEqual([r[0].value for r in ws.iter_rows(min_row=2)], [self.r.reference])

    def test_feuille_resume(self):
        _, wb = self.workbook(accident_type="OTHER")
        text = "\n".join(str(c.value) for row in wb["Résumé"].iter_rows() for c in row if c.value is not None)
        self.assertIn("ACCIMAP : export des signalements", text)
        self.assertIn("Type d'accident", text)
        self.assertIn("Autre", text)
        self.assertIn("Aucun numéro de téléphone", text)
        self.assertIn("non officielle", text)
        labels = {row[0].value: row[1].value for row in wb["Résumé"].iter_rows(min_row=1) if row[0].value}
        self.assertEqual(labels["Total"], 2)

    def test_aucun_numero_de_telephone(self):
        response, wb = self.workbook()
        text = "\n".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None)
        self.assertIsNone(PHONE_PATTERN.search(text))
        self.assertNotIn(self.owner.phone_number[4:], text)

    @override_settings(EXPORT_MAX_ROWS=2)
    def test_trop_de_lignes(self):
        self.assertContains(self.client.get(self.url, follow=True), "trop pour ce format")


class PdfExportTests(ExportDataMixin, PersonasTestCase):
    def setUp(self):
        super().setUp()
        self.make_data()
        self.login("staff")
        self.url = reverse("exports:pdf")

    def test_document_valide_avec_titre_et_pagination(self):
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))
        text = pdf_text(response)
        for expected in ["ACCIMAP", "Export des signalements", "Filtres appliqués", "Aucun filtre", "Résumé",
                         "Signalements (3)", "Page 1 sur 1", "Données collaboratives, non officielles", "non officielle"]:
            self.assertIn(expected, text)
        for report in (self.v, self.p, self.r):
            self.assertIn(report.reference, text)                 # la référence tient sur UNE ligne de cellule
            self.assertIn(report.reference, squash(text))

    def test_filtres_affiches_et_appliques(self):
        text = pdf_text(self.client.get(self.url, {"status": "VERIFIED", "mode": "identified"}))
        self.assertIn("Statut : Vérifié", text)
        self.assertIn("Mode : Identifiés", text)
        self.assertIn(self.v.reference, text)
        self.assertNotIn(self.p.reference, text)
        self.assertNotIn(self.r.reference, text)
        self.assertIn("Signalements (1)", text)

    def test_periode_affichee(self):
        today = timezone.localdate()
        text = pdf_text(self.client.get(self.url, {"date_from": (today - timedelta(days=4)).isoformat()}))
        self.assertRegex(text, r"Période \(date de l'accident\)\s*: du \d{2}/\d{2}/\d{4} au …")

    def test_resume_statistique(self):
        text = pdf_text(self.client.get(self.url))
        for expected in ["Total", "Anonymes", "Identifiés", "En attente", "Vérifiés", "Rejetés", "Blessés", "Décès", "Hors zone de couverture"]:
            self.assertIn(expected, text)

    def test_aucun_numero_de_telephone(self):
        text = pdf_text(self.client.get(self.url))
        self.assertIsNone(PHONE_PATTERN.search(text))
        self.assertIn("Aucun numéro de téléphone", text)

    def test_texte_special_ne_casse_pas_le_document(self):
        make_report(status="VERIFIED", description="a < b & <b>gras</b> ɖɛŋ \u2014 fin")
        response = self.client.get(self.url, {"q": "gras"})
        self.assertEqual(response.status_code, 200)
        text = pdf_text(response)
        self.assertIn("<b>gras</b>", text)               # affiché tel quel, jamais interprété
        self.assertIn("a < b & ", text)

    def test_tableau_sur_plusieurs_pages_avec_en_tete_repete(self):
        for _ in range(60):
            make_report(status="VERIFIED", description="ligne de remplissage pour allonger le document")
        text = pdf_text(self.client.get(self.url))
        self.assertRegex(text, r"Page 2 sur \d+")
        self.assertGreaterEqual(text.count("Référence"), 2)       # en-tête répété sur chaque page

    @override_settings(EXPORT_PDF_MAX_ROWS=2)
    def test_pdf_tronque_avec_mention_explicite(self):
        text = pdf_text(self.client.get(self.url))
        self.assertIn("Signalements (3)", text)
        self.assertIn("limité aux 2 signalements", text)
        self.assertIn("CSV ou Excel", text)

    def test_description_incluse_pour_le_staff(self):
        self.assertIn("Carrefour du marché".split()[0], pdf_text(self.client.get(self.url, {"status": "VERIFIED"})))


class ExportPermissionTests(ExportDataMixin, PersonasTestCase):
    URLS = ["exports:index", "exports:csv", "exports:xlsx", "exports:pdf"]

    def setUp(self):
        super().setUp()
        self.make_data()

    def test_matrice_des_permissions(self):
        expected = dict(anonymous=302, citizen=403, staff_no_view=403, staff=200, staff_readonly=200,
                        staff_no_delete=200, superuser=200)
        for persona, status in expected.items():
            for name in self.URLS:
                with self.subTest(persona=persona, url=name):
                    self.login(persona)
                    response = self.client.get(reverse(name))
                    self.assertEqual(response.status_code, status)
                    if status == 302:
                        self.assertIn("/login/", response["Location"])

    def test_refus_sans_aucune_donnee(self):
        for persona in ("anonymous", "citizen", "staff_no_view"):
            self.login(persona)
            for name in self.URLS[1:]:
                response = self.client.get(reverse(name))
                self.assertNotIn(self.v.reference.encode(), response.content)

    def test_get_uniquement(self):
        self.login("staff")
        for name in self.URLS[1:]:
            self.assertEqual(self.client.post(reverse(name)).status_code, 405)

    def test_les_exports_ne_sont_jamais_mis_en_cache(self):
        self.login("staff")
        for name in self.URLS:
            self.assertIn("no-store", self.client.get(reverse(name))["Cache-Control"])


class ExportPageAndAuditTests(ExportDataMixin, PersonasTestCase):
    def setUp(self):
        super().setUp()
        self.make_data()
        self.login("staff")

    def test_page_d_export(self):
        page = self.client.get(reverse("exports:index"), {"status": "VERIFIED"})
        self.assertContains(page, "Aucun numéro de téléphone")
        self.assertContains(page, '<strong>1</strong> signalement à exporter')
        self.assertContains(page, "Statut : Vérifié")
        for fmt in ("csv", "xlsx", "pdf"):
            self.assertContains(page, f'href="/dashboard/exports/{fmt}/?status=VERIFIED"')

    def test_page_d_export_sans_resultat(self):
        page = self.client.get(reverse("exports:index"), {"q": "introuvable"})
        self.assertContains(page, "rien à exporter")
        self.assertNotContains(page, 'id="download-csv"')

    def test_page_d_export_filtres_invalides(self):
        self.assertContains(self.client.get(reverse("exports:index"), {"status": "XX"}), "Filtres invalides")

    @override_settings(EXPORT_MAX_ROWS=2)
    def test_page_signale_le_depassement(self):
        page = self.client.get(reverse("exports:index"))
        self.assertContains(page, "Trop de résultats")
        self.assertNotContains(page, 'id="download-csv"')
        self.assertContains(page, 'id="download-pdf"')

    def test_liens_depuis_la_liste_et_les_statistiques_conservent_les_filtres(self):
        page = self.client.get(reverse("dashboard:report_list"), {"status": "PENDING"})
        self.assertContains(page, 'id="export-link" href="/dashboard/exports/?status=PENDING"')
        page = self.client.get(reverse("dashboard:statistics"), {"mode": "anonymous"})
        self.assertContains(page, "/dashboard/exports/?mode=anonymous")

    def test_onglet_exports_actif(self):
        html = self.client.get(reverse("exports:index")).content.decode()
        dash_nav = re.search(r'<nav class="dash-nav.*?</nav>', html, re.S).group(0)
        self.assertRegex(dash_nav, r'class="nav-link active" href="/dashboard/exports/"')
        self.assertNotRegex(dash_nav, r'class="nav-link active" href="/dashboard/"')   # « Vue d'ensemble » inactive

    def test_chaque_export_est_journalise_sans_donnee_personnelle(self):
        self.client.get(reverse("exports:csv"), {"status": "VERIFIED"})
        self.client.get(reverse("exports:xlsx"))
        self.client.get(reverse("exports:pdf"), {"mode": "anonymous"})
        entries = list(LogEntry.objects.filter(change_message__startswith="[EXPORT]").order_by("action_time"))
        self.assertEqual(len(entries), 3)
        self.assertTrue(all(e.user == self.personas["staff"] for e in entries))
        self.assertIn("CSV : 1 signalement(s). Filtres : Statut : Vérifié", entries[0].change_message)
        self.assertIn("Excel : 3 signalement(s). Filtres : aucun filtre", entries[1].change_message)
        self.assertIn("Mode : Anonymes", entries[2].change_message)
        for entry in entries:
            self.assertIsNone(PHONE_PATTERN.search(entry.change_message + entry.object_repr))

    def test_un_refus_ne_journalise_rien(self):
        self.login("citizen")
        self.client.get(reverse("exports:csv"))
        self.assertFalse(LogEntry.objects.filter(change_message__startswith="[EXPORT]").exists())
